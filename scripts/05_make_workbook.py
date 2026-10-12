"""Step 5: build the annotation workbook (annotation_sheet.xlsx). Refuses to overwrite an existing one."""
import csv, os, sys
if os.path.exists("annotation_sheet.xlsx") and "--force" not in sys.argv:
    sys.exit("annotation_sheet.xlsx exists; refusing to overwrite annotations (use --force only if you mean it).")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

F = 'Arial'
HDR = PatternFill('solid', fgColor='1F3864'); INPUT = PatternFill('solid', fgColor='FFF2CC')
EXF = PatternFill('solid', fgColor='E7E6E6'); SECF = PatternFill('solid', fgColor='D9E1F2')
thin = Side(style='thin', color='BFBFBF')

ITEMS = [  # code, label options, short question
 ('A1_split_unit', 'SAFE,RISK,UNCLEAR', 'Train/test separated by patient (or other independent unit)?'),
 ('A2_preproc_before_split', 'SAFE,RISK,UNCLEAR', 'Data-dependent preprocessing fitted on full data before split?'),
 ('A3_aug_before_split', 'SAFE,RISK,UNCLEAR,N/A', 'Augmentation / synthetic data applied before split?'),
 ('A4_test_for_selection', 'SAFE,RISK,UNCLEAR', 'Reported test set also used for tuning / early stopping / checkpoint choice?'),
 ('A5_duplicate_check', 'YES,NO', 'Duplicate / near-duplicate check reported?'),
 ('B1_external_validation', 'YES,NO', 'Tested on data from a source not used in development?'),
 ('B2_variability_reported', 'YES,NO', 'Spread across folds/seeds or confidence interval reported?'),
 ('B3_stat_test_for_claim', 'YES,NO,N/A', 'Superiority claim backed by a statistical test / non-overlapping CIs?'),
 ('B4_baseline_protocol', 'SAFE,RISK,UNCLEAR,N/A', 'Comparators evaluated under the same data and protocol?'),
 ('C1_code_available', 'YES,REQUEST,NO', 'Study code publicly available?'),
 ('C2_data_available', 'PUBLIC,RESTRICTED,NO', 'Data publicly available?'),
]
META = ['order_no', 'set', 'status', 'pmcid', 'open', 'title', 'journal', 'auto_hint', 'eligible', 'exclusion_reason']
TAIL = ['notes', 'external_note', 'minutes', 'annotated_on']
HEAD = META + [h for code, _, _ in ITEMS for h in (code, code[:2] + '_evidence')] + TAIL
col = {h: i + 1 for i, h in enumerate(HEAD)}
label_cols = [L(col[c]) for c, _, _ in ITEMS]

rows = list(csv.DictReader(open('data/processed/annotation_order.csv', encoding='utf-8')))
FIRST, LAST = 3, 2 + len(rows)

wb = Workbook()
# ---------------- Annotation sheet ----------------
ws = wb.active; ws.title = 'Annotation'
for h, c in col.items():
    cell = ws.cell(1, c, h); cell.font = Font(name=F, bold=True, color='FFFFFF'); cell.fill = HDR
    cell.alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[1].height = 32

example = {'order_no': 'EX', 'set': 'EXAMPLE', 'status': 'example', 'pmcid': 'PMC0000000', 'open': '-',
  'title': 'EXAMPLE (fictional paper, not in the corpus): ResNet-50 classification of chest radiographs',
  'journal': '-', 'auto_hint': '', 'eligible': 'Y', 'exclusion_reason': '',
  'A1_split_unit': 'RISK', 'A1_evidence': '1,200 radiographs from 400 patients || images were randomly split into 80% training and 20% testing',
  'A2_preproc_before_split': 'SAFE', 'A2_evidence': 'images were resized to 224 x 224 and scaled to [0, 1]',
  'A3_aug_before_split': 'RISK', 'A3_evidence': 'after augmentation, the dataset of 6,000 images was split',
  'A4_test_for_selection': 'UNCLEAR', 'A4_evidence': 'ABSENT', 'A5_duplicate_check': 'NO', 'A5_evidence': 'ABSENT',
  'B1_external_validation': 'NO', 'B1_evidence': 'ABSENT', 'B2_variability_reported': 'NO', 'B2_evidence': 'ABSENT',
  'B3_stat_test_for_claim': 'NO', 'B3_evidence': 'our model outperformed all compared methods',
  'B4_baseline_protocol': 'RISK', 'B4_evidence': 'Table 5 compares our accuracy with the values reported in previous studies',
  'C1_code_available': 'YES', 'C1_evidence': 'code is available at https://github.com/example/repo',
  'C2_data_available': 'PUBLIC', 'C2_evidence': 'the dataset is publicly available on Kaggle',
  'notes': 'headline result = 3-class experiment', 'external_note': '', 'minutes': 9, 'annotated_on': '2026-09-24'}
for h, v in example.items():
    c = ws.cell(2, col[h], v); c.font = Font(name=F, italic=True, color='595959'); c.fill = EXF

for r, row in enumerate(rows, start=FIRST):
    I = L(col['eligible'])
    ws.cell(r, col['order_no'], int(row['order_no']))
    ws.cell(r, col['set'], f'=IF({I}{r}="Y",IF(COUNTIF({I}${FIRST}:{I}{r},"Y")<=20,"pilot",'
                           f'IF(COUNTIF({I}${FIRST}:{I}{r},"Y")<=40,"dev",'
                           f'IF(COUNTIF({I}${FIRST}:{I}{r},"Y")<=200,"test","extra"))),"")')
    labs = ','.join(f'{c}{r}' for c in label_cols)
    ws.cell(r, col['status'], f'=IF({I}{r}="Y",IF(COUNTA({labs})={len(ITEMS)},"done","in progress"),IF({I}{r}="N","excluded",""))')
    ws.cell(r, col['pmcid'], row['pmcid'])
    link = ws.cell(r, col['open'], 'Europe PMC'); link.hyperlink = f"https://europepmc.org/article/PMC/{row['pmcid']}"
    link.font = Font(name=F, color='0563C1', underline='single')
    ws.cell(r, col['title'], row['title']); ws.cell(r, col['journal'], row['journal'])
    ws.cell(r, col['auto_hint'], row['auto_hint'])
    for h in HEAD[col['eligible'] - 1:]:
        ws.cell(r, col[h]).fill = INPUT
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(r, c)
        if cell.font.name != F or c != col['open']:
            if c != col['open']: cell.font = Font(name=F)
        cell.alignment = Alignment(vertical='top', wrap_text=HEAD[c - 1].endswith('evidence') or HEAD[c - 1] in ('title', 'notes', 'external_note'))
        cell.border = Border(bottom=thin)

widths = {'order_no': 8, 'set': 8, 'status': 11, 'pmcid': 13, 'open': 11, 'title': 48, 'journal': 22, 'auto_hint': 22,
          'eligible': 8, 'exclusion_reason': 18, 'notes': 30, 'external_note': 24, 'minutes': 8, 'annotated_on': 12}
for h, c in col.items():
    ws.column_dimensions[L(c)].width = widths.get(h, 30 if h.endswith('evidence') else 15)
ws.freeze_panes = 'I3'
ws.auto_filter.ref = f'A1:{L(len(HEAD))}{LAST}'

def dv(options, cols):
    v = DataValidation(type='list', formula1=f'"{options}"', allow_blank=True, showErrorMessage=True,
                       errorTitle='Invalid label', error=f'Choose one of: {options}')
    ws.add_data_validation(v)
    for c in cols: v.add(f'{c}{FIRST}:{c}{LAST}')
dv('Y,N', [L(col['eligible'])])
for code, opts, _ in ITEMS: dv(opts, [L(col[code])])
rng = f'{label_cols[0]}{FIRST}:{L(col["C2_data_available"])}{LAST}'
for val, color in (('RISK', 'F8CBAD'), ('UNCLEAR', 'FFE699'), ('SAFE', 'C6EFCE')):
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=color)))

# ---------------- Start here sheet ----------------
s = wb.create_sheet('Start here', 0)
s.column_dimensions['A'].width = 34; s.column_dimensions['B'].width = 90
def put(r, a, b='', bold=False, fill=None, size=10):
    ca = s.cell(r, 1, a); cb = s.cell(r, 2, b)
    for c in (ca, cb):
        c.font = Font(name=F, bold=bold, size=size); c.alignment = Alignment(wrap_text=True, vertical='top')
        if fill: c.fill = fill
put(1, 'Methodology-risk gold standard — annotation workbook', '', bold=True, size=13)
put(2, 'Codebook', 'docs/codebook_v1.0.md (full decision rules; this workbook only summarises them)')
put(4, 'WHAT YOU EDIT', '', bold=True, fill=SECF)
put(5, 'Yellow cells (sheet "Annotation")', 'eligible, exclusion_reason, the 11 label columns (drop-downs), their evidence columns, notes, external_note, minutes, annotated_on.')
put(6, 'Do not edit', 'order_no, set, status (formulas), pmcid, title, journal, auto_hint. Row 2 is a fictional EXAMPLE showing the expected format.')
put(8, 'WORKFLOW', '', bold=True, fill=SECF)
steps = ['Work strictly top-down in order_no. Open the paper via the "Europe PMC" link.',
         'Decide eligibility (codebook §3). If N, give a reason and move on; the next eligible paper takes its place automatically.',
         'Label all 11 items from the article text only. Copy verbatim evidence (ABSENT when nothing supports a label).',
         'Log minutes and date. "set" fills itself: first 20 eligible = pilot, 21-40 = dev, 41-200 = test.',
         'After the 20 pilot papers: stop, send the file back for codebook revision (v1.1) before continuing.',
         'Never look at any model prediction for a paper before its labels are final.']
for i, t in enumerate(steps): put(9 + i, f'Step {i + 1}', t)
put(16, 'LABEL MEANINGS', '', bold=True, fill=SECF)
put(17, 'SAFE / RISK / UNCLEAR / N/A', 'SAFE = the article states the safe practice. RISK = it states (or its numbers show) the risky practice. UNCLEAR = cannot tell from the text. N/A = the item does not apply.')
put(18, 'YES / NO (and variants)', 'Absence of a statement is always NO (reporting items) or UNCLEAR (risk items), never SAFE.')
put(20, 'PROGRESS', '', bold=True, fill=SECF)
A = "Annotation!"
prog = [('Papers screened', f'=COUNTIF({A}I{FIRST}:I{LAST},"Y")+COUNTIF({A}I{FIRST}:I{LAST},"N")'),
        ('Eligible', f'=COUNTIF({A}I{FIRST}:I{LAST},"Y")'),
        ('Excluded', f'=COUNTIF({A}I{FIRST}:I{LAST},"N")'),
        ('Pilot papers done (target 20)', f'=COUNTIFS({A}B{FIRST}:B{LAST},"pilot",{A}C{FIRST}:C{LAST},"done")'),
        ('Dev papers done (target 20 more)', f'=COUNTIFS({A}B{FIRST}:B{LAST},"dev",{A}C{FIRST}:C{LAST},"done")'),
        ('Test papers done (target 160)', f'=COUNTIFS({A}B{FIRST}:B{LAST},"test",{A}C{FIRST}:C{LAST},"done")'),
        ('Mean minutes per completed paper', f'=IFERROR(AVERAGEIF({A}C{FIRST}:C{LAST},"done",{A}{L(col["minutes"])}{FIRST}:{L(col["minutes"])}{LAST}),"-")')]
for i, (a, f) in enumerate(prog):
    put(21 + i, a); s.cell(21 + i, 2, f).font = Font(name=F, bold=True); s.cell(21 + i, 2).alignment = Alignment(horizontal='left')
put(29, 'Source of the paper list', 'Europe PMC, 2025, CC BY / CC0, deep-learning image classification (542 papers). Seeded random order (seed 20260923). See data/processed/frame_log.json.')

# ---------------- Codebook quick reference ----------------
q = wb.create_sheet('Codebook (short)')
for j, h in enumerate(['Item', 'Question', 'Labels'], 1):
    c = q.cell(1, j, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HDR
for i, (code, opts, ques) in enumerate(ITEMS, 2):
    for j, v in enumerate([code, ques, opts.replace(',', ' / ')], 1):
        c = q.cell(i, j, v); c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical='top')
q.column_dimensions['A'].width = 26; q.column_dimensions['B'].width = 70; q.column_dimensions['C'].width = 30
q.cell(len(ITEMS) + 3, 1, 'Full decision rules and edge cases: docs/codebook_v1.0.md').font = Font(name=F, italic=True)

wb.save('annotation_sheet.xlsx'); print('saved', len(rows), 'papers,', len(HEAD), 'columns')
