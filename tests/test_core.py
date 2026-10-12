import json, os, random, shutil, tempfile
import pytest
from mrscreen import stats, corpus, retrieval
from mrscreen.items import ITEMS, ITEM_BY_CODE
from mrscreen.verify import Verifier, apply_policy


def test_kappa_and_f1():
    assert stats.cohen_kappa(list("aabb"), list("aabb")) == 1.0
    # classic 2x2 example: po=0.7, pe=0.5 -> kappa=0.4
    y = ["y"] * 5 + ["n"] * 5
    p = ["y"] * 4 + ["n"] + ["y"] * 2 + ["n"] * 3
    assert abs(stats.cohen_kappa(y, p) - 0.4) < 1e-9
    assert 0 <= stats.macro_f1(y, p) <= 1


def test_mcnemar_holm_wilson():
    a = [True] * 10 + [False] * 10
    b = [False] * 10 + [True] * 10
    assert stats.mcnemar_exact(a, b)["p"] == 1.0
    r = stats.mcnemar_exact([True] * 12, [False] * 12)
    assert r["b"] == 12 and r["p"] < 0.001
    adj = stats.holm({"x": 0.01, "y": 0.04, "z": 0.03})
    assert adj["x"] == pytest.approx(0.03) and adj["y"] >= adj["z"] >= adj["x"]
    lo, hi = stats.wilson(5, 10)
    assert 0 < lo < 0.5 < hi < 1


def test_verifier_levels_and_policy():
    text = "The images were randomly split into training (80%) and test (20%) sets [12]. Code is on GitHub."
    v = Verifier(text)
    assert v.check("images were randomly split into training (80%) and test (20%) sets") == "exact"
    assert v.check("The images were randomly split into training (80 %) and test (20%) sets") in ("exact", "fuzzy")
    assert v.check("patients were split at the patient level before training") == "none"
    it = ITEM_BY_CODE["A1_split_unit"]
    assert apply_policy(it, "RISK", ["none"], "verified") == "UNCLEAR"
    assert apply_policy(it, "RISK", ["fuzzy"], "verified") == "RISK"
    assert apply_policy(it, "RISK", ["none"], "raw") == "RISK"
    assert apply_policy(ITEM_BY_CODE["A3_aug_before_split"], "N/A", [], "verified") == "N/A"


def test_retrieval_order_and_budget():
    pid = corpus.select("sandbox:1")[0]
    psgs = corpus.passages(corpus.load_paper(pid))
    for it in ITEMS:
        sel = retrieval.retrieve(it, psgs, 1200)
        assert [p.pid for p in sel] == sorted(p.pid for p in sel)
        assert sum(corpus.approx_tokens(p.text) + 8 for p in sel) <= 1200


def test_sets_rule():
    elig = {i: "Y" for i in range(1, 230)}
    elig[3] = "N"
    s = corpus.assign_sets(elig)
    assert s[21] == "pilot" and s[22] == "dev" and s[41] == "dev" and s[42] == "test" and s[121] == "test" and s[122] == "extra" and 3 not in s


def synthetic_workbook(tmp):
    """Copy of the real workbook where 40 SANDBOX papers get random labels.
    SYNTHETIC - for testing the code path only; never a result."""
    from openpyxl import load_workbook
    src = os.path.join(corpus.ROOT, "templates", "annotation_sheet_blank.xlsx")
    dst = os.path.join(tmp, "synthetic.xlsx")
    wb = load_workbook(src)
    ws = wb["Annotation"]
    head = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(head)}
    rng = random.Random(0)
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, col["order_no"]).value and 301 <= ws.cell(r, col["order_no"]).value <= 340:
            ws.cell(r, col["eligible"], "Y")
            for it in ITEMS:
                ws.cell(r, col[it.code], rng.choice(it.labels))
    wb.save(dst)
    return dst


def test_end_to_end_mock_and_rules(tmp_path):
    from mrscreen import run, evaluate
    wb = synthetic_workbook(str(tmp_path))
    out = str(tmp_path / "runs")
    d1 = run.main(["--backend", "mock", "--set", "dev", "--workbook", wb, "--out", out, "--run-name", "mock"])
    d2 = run.main(["--backend", "rules", "--set", "dev", "--workbook", wb, "--out", out, "--run-name", "rules"])
    n = sum(1 for _ in open(os.path.join(d1, "predictions.jsonl"), encoding="utf-8"))
    assert n == 40 * len(ITEMS)
    # resume: second call adds nothing
    run.main(["--backend", "rules", "--set", "dev", "--workbook", wb, "--out", out, "--run-name", "rules"])
    assert sum(1 for _ in open(os.path.join(d2, "predictions.jsonl"), encoding="utf-8")) == 40 * len(ITEMS)
    evaluate.main(["--set", "dev", "--workbook", wb, "--runs", d1, d2, "--boot", "200", "--out", str(tmp_path / "res")])
    for f in ("metrics.csv", "efficiency.csv", "comparisons.csv", "prevalence.csv", "summary.md"):
        assert os.path.exists(tmp_path / "res" / f)


def test_test_set_guard(tmp_path):
    from mrscreen import run
    with pytest.raises(SystemExit):
        run.main(["--backend", "rules", "--set", "test", "--out", str(tmp_path)])


def test_check_annotations_flags_wrong_quotes(tmp_path):
    """SYNTHETIC: sandbox papers 301-302 get one real quote, one foreign quote, one invented quote."""
    from openpyxl import load_workbook
    from mrscreen import check_annotations
    src = os.path.join(corpus.ROOT, "templates", "annotation_sheet_blank.xlsx")
    wb = load_workbook(src); ws = wb["Annotation"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    rows = {ws.cell(r, col["order_no"]).value: r for r in range(3, ws.max_row + 1)}
    p1, p2 = ws.cell(rows[301], col["pmcid"]).value, ws.cell(rows[302], col["pmcid"]).value
    s1 = [s for s in corpus.load_paper(p1)["sections"] if len(s["text"]) > 200][0]["text"][:150]
    s2 = [s for s in corpus.load_paper(p2)["sections"] if len(s["text"]) > 200][0]["text"][:150]
    for o in (301, 302):
        r = rows[o]; ws.cell(r, col["eligible"], "Y")
        for it in ITEMS:
            ws.cell(r, col[it.code], it.absent_label)
    ws.cell(rows[301], col["A1_split_unit"], "SAFE"); ws.cell(rows[301], col["A1_evidence"], s1)          # correct
    ws.cell(rows[302], col["A1_split_unit"], "SAFE"); ws.cell(rows[302], col["A1_evidence"], s1)          # other paper
    ws.cell(rows[302], col["B1_external_validation"], "YES"); ws.cell(rows[302], col["B1_evidence"], "An invented sentence that appears in no paper at all.")
    ws.cell(rows[302], col["B2_variability_reported"], "YES")                                              # no evidence
    path = str(tmp_path / "wb.xlsx"); wb.save(path)
    _, problems, _, _ = check_annotations.check(path, 301, 302)
    kinds = {(g["order_no"], item, kind) for g, item, kind, _ in problems}
    assert (302, "A1_split_unit", "IN_OTHER_PAPER") in kinds
    assert (302, "B1_external_validation", "NOT_IN_PAPER") in kinds
    assert (302, "B2_variability_reported", "NO_EVIDENCE") in kinds
    assert not any(o == 301 and item == "A1_split_unit" for o, item, _ in kinds)


def test_decomposed_mapping_follows_codebook():
    from mrscreen.decomposed import map_label, schema, FACTS
    assert map_label("A1_split_unit", {"split_level": "patient_or_subject_or_case"}) == "SAFE"
    assert map_label("A1_split_unit", {"split_level": "image_slice_patch_or_frame", "several_images_per_patient": "yes",
                                       "test_is_separate_cohort": "no", "one_image_per_patient": "not_stated"}) == "RISK"
    assert map_label("A1_split_unit", {"split_level": "not_stated", "test_is_separate_cohort": "no"}) == "UNCLEAR"
    assert map_label("A2_preproc_before_split", {"data_dependent_steps": [], "fitted_on": "not_stated"}) == "N/A"
    assert map_label("A2_preproc_before_split", {"data_dependent_steps": ["SMOTE"], "fitted_on": "full_dataset_before_split"}) == "RISK"
    assert map_label("A2_preproc_before_split", {"data_dependent_steps": ["PCA"], "fitted_on": "not_stated"}) == "UNCLEAR"
    assert map_label("A3_aug_before_split", {"augmentation_or_synthetic_used": "yes", "timing": "before_the_split",
                                             "split_grouped_by_source": "yes"}) == "SAFE"
    assert map_label("A3_aug_before_split", {"augmentation_or_synthetic_used": "no"}) == "N/A"
    assert map_label("A4_test_for_selection", {"design": "train_validation_test", "reported_set_used_for_decisions": "no",
                                               "best_model_chosen_on_reported_set": "yes"}) == "RISK"
    assert map_label("A4_test_for_selection", {"design": "train_test_only", "reported_set_used_for_decisions": "not_stated",
                                               "best_model_chosen_on_reported_set": "no"}) == "UNCLEAR"
    assert map_label("B3_stat_test_for_claim", {"superiority_claim": "no", "statistical_support": "yes"}) == "N/A"
    assert map_label("B4_baseline_protocol", {"compares_with_other_methods_or_readers": "yes",
                                              "comparator_source": "numbers_taken_from_other_papers"}) == "RISK"
    assert map_label("C1_code_available", {"code_link": "https://github.com/x/y", "code_on_request": "no"}) == "YES"
    assert map_label("C1_code_available", {"code_link": "", "code_on_request": "yes"}) == "REQUEST"
    assert map_label("C2_data_available", {"all_data_public_and_named": "no", "available_on_request_or_restricted": "yes"}) == "RESTRICTED"
    for code in FACTS:     # every schema lists evidence first and all fact fields
        s = schema(code)
        assert list(s["properties"])[0] == "evidence" and len(s["required"]) == len(FACTS[code]) + 1


def decomposed_version():
    from mrscreen import decomposed
    return decomposed.PROMPT_VERSION


def test_decomposed_mock_run(tmp_path):
    from mrscreen import run
    d = run.main(["--backend", "mock", "--prompt", "decomposed", "--set", "sandbox:2", "--out", str(tmp_path), "--run-name", "p2"])
    rows = [json.loads(l) for l in open(os.path.join(d, "predictions.jsonl"), encoding="utf-8")]
    assert len(rows) == 2 * len(ITEMS) and all(r["label_raw"] in ITEM_BY_CODE[r["item"]].labels for r in rows)
    meta = json.load(open(os.path.join(d, "run_meta.json"), encoding="utf-8"))
    assert meta["prompt_version"].endswith(decomposed_version())


def test_gated_policy():
    it = ITEM_BY_CODE["B2_variability_reported"]
    assert apply_policy(it, "YES", ["exact"], "gated", ["Five-fold cross-validation was used for all models."]) == "NO"
    assert apply_policy(it, "YES", ["exact"], "gated", ["The mean AUC was 0.91 ± 0.02 across folds."]) == "YES"
    assert apply_policy(it, "YES", ["exact"], "verified", ["Five-fold cross-validation was used for all models."]) == "YES"


def test_backend_down_stops_and_resume_reruns_failed(tmp_path, monkeypatch):
    """A run whose backend is unreachable must stop without writing rows; failed papers are re-run on resume."""
    from mrscreen import run, backends
    out = str(tmp_path)
    class Down:
        def chat(self, *a, **k): return backends.Reply(None, error="connection refused")
        def describe(self): return {"backend": "down"}
    monkeypatch.setattr(run, "make_backend", lambda *a, **k: Down())
    with pytest.raises(SystemExit):
        run.main(["--backend", "mock", "--set", "sandbox:2", "--out", out, "--run-name", "r"])
    p = os.path.join(out, "r", "predictions.jsonl")
    assert not os.path.exists(p) or sum(1 for _ in open(p, encoding="utf-8")) == 0
    # simulate an old file where paper 1 failed and paper 2 succeeded, then resume with a working backend
    ids = corpus.select("sandbox:2")
    with open(p, "w", encoding="utf-8") as f:
        for pid, err in ((ids[0], ["refused"]), (ids[1], [])):
            for it in ITEMS:
                f.write(json.dumps({"pmcid": pid, "item": it.code, "errors": err}) + "\n")
    monkeypatch.setattr(run, "make_backend", backends.make_backend)
    run.main(["--backend", "mock", "--set", "sandbox:2", "--out", out, "--run-name", "r"])
    rows = [json.loads(l) for l in open(p, encoding="utf-8")]
    assert len(rows) == 2 * len(ITEMS) and not any(r["errors"] for r in rows)


def test_energy_meter_handles_na_power(monkeypatch):
    from mrscreen import energy
    m = energy.GpuMeter()
    m.samples = [(0.0, None, 900.0), (1.0, None, 5000.0), (2.0, None, 5100.0)]
    m.t0, m.t1 = 0.0, 2.0
    s = m.summary()
    assert s["energy_j"] is None and s["peak_vram_mib"] == 5100.0
    m.samples = [(0.0, 50.0, 900.0), (1.0, 70.0, 5000.0), (2.0, 70.0, 5100.0)]
    assert abs(m.summary()["energy_j"] - 130.0) < 1e-9


def test_five_item_run_and_check(tmp_path):
    from mrscreen import run
    from mrscreen.items import TEST_ITEMS
    d = run.main(["--backend", "rules", "--items", "test5", "--set", "sandbox:2", "--out", str(tmp_path), "--run-name", "r5"])
    rows = [json.loads(l) for l in open(os.path.join(d, "predictions.jsonl"), encoding="utf-8")]
    assert len(rows) == 2 * 5 and {r["item"] for r in rows} == set(TEST_ITEMS)
